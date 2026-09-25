"""Turn an uploaded file into pages of text blocks and tables (guide 7.1 steps 2-7).

No database code here, so every parser can be unit-tested on its own. Coordinates:
- PDF: PDF points (1/72 inch) on that page, the same space PyMuPDF uses to draw highlights.
- Images: pixels of the original image.
- Excel / CSV / DOCX: no page geometry, bbox is None.
"""

from __future__ import annotations

from collections.abc import Callable, Iterator
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from PIL import Image, ImageSequence

from worker.ocr import OcrResult

BBox = list[float]  # [x0, y0, x1, y1]
OcrFn = Callable[[Image.Image], OcrResult]

SCANNED_TEXT_CHARS = 50  # guide 7.1 step 2: fewer characters than this -> treat as scanned
OCR_DPI = 300


@dataclass
class Block:
    text: str
    bbox: BBox | None = None


@dataclass
class Table:
    rows: list[list[str]]
    bbox: BBox | None = None
    title: str | None = None


@dataclass
class PageResult:
    page_no: int  # 1-based; sheet number for spreadsheets
    has_text_layer: bool
    blocks: list[Block] = field(default_factory=list)
    tables: list[Table] = field(default_factory=list)
    ocr_used: bool = False
    ocr_confidence: float | None = None


def page_count(path: Path, mime: str) -> int:
    if mime == "application/pdf":
        with pymupdf.open(path) as doc:
            return doc.page_count
    if mime.startswith("image/"):
        with Image.open(path) as img:
            return getattr(img, "n_frames", 1)
    if mime in (SHEET_MIMES | {"text/csv"}):
        return len(_read_sheets(path, mime))
    return 1


def parse(path: Path, mime: str, ocr: OcrFn) -> Iterator[PageResult]:
    if mime == "application/pdf":
        yield from parse_pdf(path, ocr)
    elif mime.startswith("image/"):
        yield from parse_image(path, ocr)
    elif mime in SHEET_MIMES or mime == "text/csv":
        yield from parse_sheets(path, mime)
    elif mime == DOCX_MIME:
        yield from parse_docx(path)
    else:
        raise ValueError(f"No parser for {mime}")


# --- PDF -----------------------------------------------------------------------------------


def parse_pdf(path: Path, ocr: OcrFn) -> Iterator[PageResult]:
    with pymupdf.open(path) as doc:
        for index, page in enumerate(doc):
            if len(page.get_text("text").strip()) < SCANNED_TEXT_CHARS:
                yield _ocr_pdf_page(page, index + 1, ocr)
            else:
                yield _digital_pdf_page(page, index + 1)


def _digital_pdf_page(page: pymupdf.Page, page_no: int) -> PageResult:
    blocks = [
        Block(text.strip(), [round(v, 2) for v in (x0, y0, x1, y1)])
        for x0, y0, x1, y1, text, _no, kind in page.get_text("blocks", sort=True)
        if kind == 0 and text.strip()  # kind 0 = text, 1 = image
    ]
    return PageResult(page_no, has_text_layer=True, blocks=blocks, tables=_pdf_tables(page))


def _pdf_tables(page: pymupdf.Page) -> list[Table]:
    try:
        found = page.find_tables()
    except Exception:  # table detection is best effort; never fail the page because of it
        return []
    tables = []
    for t in found.tables:
        rows = [[_cell(c) for c in row] for row in t.extract()]
        if len(rows) >= 2 and any(any(c for c in row) for row in rows):
            tables.append(Table(rows, [round(v, 2) for v in t.bbox]))
    return tables


def _ocr_pdf_page(page: pymupdf.Page, page_no: int, ocr: OcrFn) -> PageResult:
    pix = page.get_pixmap(dpi=OCR_DPI)
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    result = ocr(image)
    scale = 72 / OCR_DPI  # OCR pixels -> PDF points
    blocks = [Block(line.text, [round(v * scale, 2) for v in line.bbox]) for line in result.lines]
    return PageResult(
        page_no,
        has_text_layer=False,
        blocks=blocks,
        ocr_used=True,
        ocr_confidence=result.confidence,
    )


# --- Images ----------------------------------------------------------------------------------


def parse_image(path: Path, ocr: OcrFn) -> Iterator[PageResult]:
    with Image.open(path) as img:
        for index, frame in enumerate(ImageSequence.Iterator(img)):  # multi-page TIFF
            result = ocr(frame.convert("RGB"))
            yield PageResult(
                index + 1,
                has_text_layer=False,
                blocks=[Block(line.text, line.bbox) for line in result.lines],
                ocr_used=True,
                ocr_confidence=result.confidence,
            )


# --- Excel / CSV -----------------------------------------------------------------------------

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS_MIME = "application/vnd.ms-excel"
SHEET_MIMES = {XLSX_MIME, XLS_MIME}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _read_sheets(path: Path, mime: str) -> dict[str, list[list[str]]]:
    import pandas as pd

    if mime == "text/csv":
        try:
            frames = {"CSV": pd.read_csv(path, header=None, dtype=str, encoding="utf-8-sig")}
        except UnicodeDecodeError:
            frames = {"CSV": pd.read_csv(path, header=None, dtype=str, encoding="cp1252")}
    else:
        frames = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
    # Cells are kept exactly as text; numbers are never re-typed (guide 7.2).
    return {
        name: [[_cell(v) for v in row] for row in df.itertuples(index=False)]
        for name, df in frames.items()
    }


def parse_sheets(path: Path, mime: str) -> Iterator[PageResult]:
    """Every sheet is one table (guide 7.1 step 7), plus row text for keyword search."""
    for index, (name, rows) in enumerate(_read_sheets(path, mime).items()):
        rows = [r for r in rows if any(r)]
        blocks = [Block(" | ".join(c for c in row if c)) for row in rows]
        tables = [Table(rows, title=str(name))] if rows else []
        yield PageResult(index + 1, has_text_layer=True, blocks=blocks, tables=tables)


# --- Word --------------------------------------------------------------------------------------


def parse_docx(path: Path) -> Iterator[PageResult]:
    import docx  # python-docx

    document = docx.Document(str(path))
    blocks = [Block(p.text.strip()) for p in document.paragraphs if p.text.strip()]
    tables = []
    for t in document.tables:
        rows = [[_cell(c.text) for c in row.cells] for row in t.rows]
        if rows:
            tables.append(Table(rows))
    # A .docx has no fixed pages; everything is page 1.
    yield PageResult(1, has_text_layer=True, blocks=blocks, tables=tables)


def _cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    return "" if text.lower() == "nan" else text.strip()
