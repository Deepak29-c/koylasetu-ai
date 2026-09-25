"""Parser tests on small generated files. OCR is faked, so no Tesseract is needed here."""

import csv
import shutil

import pymupdf
import pytest
from PIL import Image

from worker import parse
from worker.embed import chunk_blocks
from worker.ocr import OcrLine, OcrResult, tesseract_lang
from worker.parse import Block

PARA = (
    "Coal production of CIL during the year was reported subsidiary-wise in the annual "
    "report. This paragraph is long enough to count as a real text layer on the page."
)


def fake_ocr(image: Image.Image) -> OcrResult:
    # One line covering the top-left 100x50 px of the OCR image.
    return OcrResult([OcrLine("SCANNED TEXT", [0, 0, 100, 50], 0.9)], 0.9)


@pytest.fixture
def sample_pdf(tmp_path):
    """Page 1: digital text + a 3x3 ruled table. Page 2: image only (like a scan)."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 150), PARA, fontsize=11)
    x0, y0, cw, rh = 50, 200, 150, 25
    for r in range(4):
        page.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
    for c in range(4):
        page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
    cells = [["Subsidiary", "2023-24", "2024-25"], ["MCL", "10", "11"], ["SECL", "20", "21"]]
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            page.insert_text((x0 + c * cw + 5, y0 + r * rh + 17), value, fontsize=10)

    scan = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), False)
    pix.clear_with(200)
    scan.insert_image(pymupdf.Rect(50, 50, 450, 250), pixmap=pix)

    path = tmp_path / "sample.pdf"
    doc.save(path)
    return path


def test_pdf_digital_and_scanned_pages(sample_pdf):
    assert parse.page_count(sample_pdf, "application/pdf") == 2
    p1, p2 = list(parse.parse(sample_pdf, "application/pdf", fake_ocr))

    assert p1.has_text_layer and not p1.ocr_used
    assert any("Coal production" in b.text for b in p1.blocks)
    assert all(b.bbox and b.bbox[2] > b.bbox[0] for b in p1.blocks)

    assert not p2.has_text_layer and p2.ocr_used and p2.ocr_confidence == 0.9
    assert p2.blocks[0].text == "SCANNED TEXT"
    # 100x50 px at 300 DPI -> 24x12 PDF points
    assert p2.blocks[0].bbox == [0.0, 0.0, 24.0, 12.0]


def test_pdf_table_cells_are_copied_as_text(sample_pdf):
    p1 = next(parse.parse(sample_pdf, "application/pdf", fake_ocr))
    assert len(p1.tables) == 1
    rows = p1.tables[0].rows
    assert rows[0] == ["Subsidiary", "2023-24", "2024-25"]
    assert rows[1] == ["MCL", "10", "11"]
    assert p1.tables[0].bbox is not None


def test_csv_sheet_is_one_table_with_exact_cell_text(tmp_path):
    path = tmp_path / "prod.csv"
    with path.open("w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows([["Subsidiary", "FY 2024-25"], ["MCL", "1,234.50"], ["NCL", ""]])
    (page,) = parse.parse(path, "text/csv", fake_ocr)
    assert page.tables[0].rows == [["Subsidiary", "FY 2024-25"], ["MCL", "1,234.50"], ["NCL", ""]]
    assert page.blocks == []  # searchable text comes from the table chunk


def test_xlsx_every_sheet_is_a_page(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Production"
    wb.active.append(["Subsidiary", "2024-25"])
    wb.active.append(["MCL", "225.2"])
    wb.create_sheet("Dispatch").append(["Subsidiary", "2024-25"])
    path = tmp_path / "stats.xlsx"
    wb.save(path)

    pages = list(parse.parse(path, parse.XLSX_MIME, fake_ocr))
    assert [p.tables[0].title for p in pages] == ["Production", "Dispatch"]
    assert pages[0].tables[0].rows[1] == ["MCL", "225.2"]
    assert parse.page_count(path, parse.XLSX_MIME) == 2


def test_docx_paragraphs_and_tables(tmp_path):
    import docx

    d = docx.Document()
    d.add_paragraph("Steps taken to increase production.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(1, 0).text = "Subsidiary", "CCL"
    path = tmp_path / "note.docx"
    d.save(path)

    (page,) = parse.parse(path, parse.DOCX_MIME, fake_ocr)
    assert page.blocks[0].text == "Steps taken to increase production."
    assert page.tables[0].rows[1][0] == "CCL"


def test_multipage_tiff_is_ocrd_per_frame(tmp_path):
    path = tmp_path / "scan.tiff"
    frames = [Image.new("RGB", (300, 200), "white") for _ in range(3)]
    frames[0].save(path, save_all=True, append_images=frames[1:])
    pages = list(parse.parse(path, "image/tiff", fake_ocr))
    assert [p.page_no for p in pages] == [1, 2, 3]
    assert all(p.ocr_used for p in pages)


def test_chunks_respect_word_limit_and_merge_boxes():
    blocks = [
        Block(" ".join(["w"] * 200), [0, 0, 10, 10]),
        Block(" ".join(["x"] * 100), [5, 20, 30, 40]),
    ]
    (chunk,) = chunk_blocks(blocks, max_words=350)
    assert chunk.bbox == [0, 0, 30, 40]

    three = chunk_blocks(blocks + [Block(" ".join(["y"] * 100), [0, 50, 10, 60])], max_words=350)
    assert [len(c.text.split()) for c in three] == [300, 100]


def test_long_block_is_split_and_keeps_its_box():
    chunks = chunk_blocks([Block(" ".join(["z"] * 800), [1, 2, 3, 4])], max_words=350)
    assert [len(c.text.split()) for c in chunks] == [350, 350, 100]
    assert all(c.bbox == [1, 2, 3, 4] for c in chunks)


def test_tesseract_language_from_document_language():
    assert tesseract_lang("en") == "eng"
    assert tesseract_lang("hi") == "hin+eng"
    assert tesseract_lang("mixed") == "hin+eng"


@pytest.mark.skipif(
    shutil.which("tesseract") is None, reason="Tesseract not installed (runs in Docker)"
)
def test_real_tesseract_reads_rendered_text():
    from PIL import ImageDraw, ImageFont

    from worker.ocr import tesseract_ocr

    img = Image.new("RGB", (1400, 200), "white")
    font = ImageFont.load_default(size=60)
    ImageDraw.Draw(img).text((40, 60), "COAL PRODUCTION 2024", fill="black", font=font)
    result = tesseract_ocr(img, "eng")
    assert "COAL" in " ".join(line.text for line in result.lines)
    assert result.confidence and result.confidence > 0.5


@pytest.mark.parametrize("angle", [90, 180, 270])
def test_unrotated_boxes_land_back_on_the_original_image(angle):
    """A box drawn on an image, found on the rotated copy, must map back to where it was."""
    from worker.ocr import _unrotate_bbox

    w, h = 400, 200
    original = Image.new("L", (w, h), 0)
    original.paste(255, (50, 20, 120, 60))  # white box at x 50-120, y 20-60
    rotated = original.rotate(-angle, expand=True)  # clockwise, like ocr_upright
    found = list(rotated.getbbox())  # where the box ended up after rotation
    assert _unrotate_bbox(found, angle, w, h) == pytest.approx([50, 20, 120, 60], abs=1)


@pytest.mark.skipif(
    shutil.which("tesseract") is None, reason="Tesseract not installed (runs in Docker)"
)
def test_sideways_page_is_read_after_rotation():
    from PIL import ImageDraw, ImageFont

    from worker.ocr import ocr_upright

    img = Image.new("RGB", (1600, 500), "white")
    font = ImageFont.load_default(size=50)
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(
        ["INDIAN COAL RESOURCES", "MEASURED INDICATED INFERRED", "STATE WISE"]
    ):
        draw.text((40, 40 + i * 140), line, fill="black", font=font)
    sideways = img.rotate(90, expand=True)  # like a landscape chart in a portrait PDF
    result = ocr_upright(sideways, "eng")
    text = " ".join(line.text for line in result.lines)
    assert "RESOURCES" in text and "MEASURED" in text
    assert all(0 <= b <= max(sideways.size) for line in result.lines for b in line.bbox)


def test_pdf_table_title_cell_boxes_and_running_header(tmp_path):
    """The text just above a table is its title (with the unit); every cell keeps its box."""
    doc = pymupdf.open()
    for n in range(2):
        page = doc.new_page()
        page.insert_text((50, 30), f"Annual Report 2024-25   page {n + 1}", fontsize=9)
        page.insert_textbox(pymupdf.Rect(50, 60, 550, 160), PARA, fontsize=11)
        page.insert_text((50, 180), "Coal Production (in MT)", fontsize=11)
        x0, y0, cw, rh = 50, 200, 150, 25
        for r in range(4):
            page.draw_line((x0, y0 + r * rh), (x0 + 2 * cw, y0 + r * rh))
        for c in range(3):
            page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 3 * rh))
        for r, row in enumerate([["Subsidiary", "2024-25"], ["MCL", "225.2"], ["NCL", "139.1"]]):
            for c, value in enumerate(row):
                page.insert_text((x0 + c * cw + 5, y0 + r * rh + 17), value, fontsize=10)
    path = tmp_path / "report.pdf"
    doc.save(path)

    p1 = next(parse.parse(path, "application/pdf", fake_ocr))
    (table,) = p1.tables
    assert table.title == "Coal Production (in MT)"
    assert len(table.cell_bboxes) == len(table.rows)
    x0, y0, x1, y1 = table.cell_bboxes[1][1]  # the "225.2" cell
    assert (x0, x1) == pytest.approx((200, 350), abs=1)
    assert (y0, y1) == pytest.approx((225, 250), abs=1)


def test_running_header_is_not_a_title(tmp_path):
    doc = pymupdf.open()
    for _ in range(2):
        page = doc.new_page()
        page.insert_text((50, 40), "Ministry of Coal", fontsize=9)
        page.insert_textbox(pymupdf.Rect(50, 100, 550, 200), PARA, fontsize=11)
        x0, y0, cw, rh = 50, 60, 150, 20  # table right under the header
        for r in range(3):
            page.draw_line((x0, y0 + r * rh), (x0 + 2 * cw, y0 + r * rh))
        for c in range(3):
            page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + 2 * rh))
        for r, row in enumerate([["State", "2024-25"], ["ODISHA", "10"]]):
            for c, value in enumerate(row):
                page.insert_text((x0 + c * cw + 5, y0 + r * rh + 14), value, fontsize=10)
    path = tmp_path / "chapter.pdf"
    doc.save(path)
    p1 = next(parse.parse(path, "application/pdf", fake_ocr))
    assert p1.tables and p1.tables[0].title is None
