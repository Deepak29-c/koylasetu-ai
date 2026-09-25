"""PQ reply builder and export (guide 10.1-10.2). The Q&A service and the model are mocked."""

import io
import json
import uuid
from datetime import date, datetime
from types import SimpleNamespace

import docx
import pymupdf

from app.llm import client
from app.models import ParliamentQuestion, Report, ReportCitation, ReportVersion
from app.models.enums import House, QuestionType, ReportStatus
from app.services import pq_reply, qa, report_export
from app.services.pq_reply import Part
from app.services.retrieval import Passage

PQ_EN = """Will the Minister of COAL be pleased to state:
(a) the coal production of Northern Coalfields Limited (NCL) in December, 2024;
(b) whether the target for the month was achieved; and
(c) if so, the details thereof?"""

PQ_HI = """क्या कोयला मंत्री यह बताने की कृपा करेंगे कि:
(क) दिसंबर, 2024 में एनसीएल का कोयला उत्पादन कितना था;
(ख) क्या उस माह का लक्ष्य प्राप्त किया गया; और
(ग) यदि हां, तो तत्संबंधी ब्यौरा क्या है?"""


def passage(ref, title="Monthly Summary Dec 2024", page=1, text="NCL produced 12.70 MT."):
    return Passage(ref, uuid.uuid4(), uuid.uuid4(), title, page, None, text)


def fact(ref, who="NCL", value="12.70"):
    return SimpleNamespace(
        ref=ref,
        fact_id=uuid.uuid4(),
        document_id=uuid.uuid4(),
        document_title="Coal Directory",
        page_no=7,
        bbox=[1, 2, 3, 4],
        who=who,
        value=value,
    )


def answer(text, citations=(), not_found=False, unmatched=(), table=None, facts=()):
    return SimpleNamespace(
        answer=text,
        citations=list(citations),
        not_found=not_found,
        unmatched_numbers=list(unmatched),
        table=table,
        facts=list(facts),
    )


def header(**kw):
    base = dict(
        house=House.LOK_SABHA,
        q_type=QuestionType.UNSTARRED,
        q_no="1234",
        answer_date=date(2025, 2, 12),
        mp_name="Shri A. B. Member",
        subject="Coal production by NCL",
        intro="",
    )
    return pq_reply.build_header(**(base | kw))


# --- splitting ----------------------------------------------------------------------------------


def test_split_english_parts():
    intro, parts = pq_reply.split_by_markers(PQ_EN)
    assert intro == "Will the Minister of COAL be pleased to state:"
    assert [p.label for p in parts] == ["a", "b", "c"]
    assert parts[0].text.startswith("the coal production of Northern Coalfields")
    assert parts[2].text == "if so, the details thereof?"


def test_split_hindi_parts():
    intro, parts = pq_reply.split_by_markers(PQ_HI)
    assert intro.startswith("क्या कोयला मंत्री")
    assert [p.label for p in parts] == ["क", "ख", "ग"]
    assert "एनसीएल" in parts[0].text


def test_split_single_line_and_out_of_sequence_markers():
    q = "state: (a) the steps taken (i) to raise output (ii) to cut imports; and (b) the result?"
    _, parts = pq_reply.split_by_markers(q)
    assert [p.label for p in parts] == ["a", "b"]
    assert "(i) to raise output (ii)" in parts[0].text  # sub-items stay inside part (a)


def test_marker_inside_word_is_not_a_part():
    _, parts = pq_reply.split_by_markers("Is coal India(a) ltd listed?")
    assert parts == []


def test_no_markers_uses_llm_split_and_checks_its_parts():
    q = "Will the Minister state the coal production in 2024 and the steps taken to raise it?"

    def llm(system, user, json_schema=None):
        assert json_schema is not None
        return json.dumps(
            {
                "parts": [
                    {"text": "the coal production in 2024"},
                    {"text": "the steps taken to raise it"},
                ]
            }
        )

    intro, parts, method = pq_reply.split_question(q, llm)
    assert method == "llm" and intro == ""
    assert [(p.label, p.text) for p in parts] == [
        ("a", "the coal production in 2024"),
        ("b", "the steps taken to raise it"),
    ]


def test_llm_split_that_invents_text_or_fails_falls_back_to_one_part():
    q = "Will the Minister state the coal production in 2024?"

    def inventing(system, user, json_schema=None):
        return json.dumps({"parts": [{"text": "total imports of lignite"}, {"text": "prices"}]})

    def down(*a, **k):
        raise client.LlmUnavailable("down")

    for llm in (inventing, down, lambda *a, **k: "not json"):
        _, parts, method = pq_reply.split_question(q, llm)
        assert method == "single"
        assert parts == [Part("", q)]


def test_markers_do_not_call_the_model():
    def boom(*a, **k):
        raise AssertionError("model called")

    assert pq_reply.split_question(PQ_EN, boom)[2] == "markers"


def test_dependent_part_carries_previous_part_and_subject():
    _, parts = pq_reply.split_by_markers(PQ_EN)
    q = pq_reply.part_query("Coal production", parts, 2)
    assert q.startswith("Coal production: whether the target")
    assert q.endswith("if so, the details thereof?")
    assert pq_reply.part_query(None, parts, 0) == parts[0].text


# --- blocks, annexures, citations ---------------------------------------------------------------


def test_block_renumbers_citations_across_parts():
    reg = pq_reply.SourceRegistry()
    shared = passage("C1")
    other = passage("C1", title="Annual Report", page=33, text="Target 12.16 MT.")
    b1 = pq_reply.build_block(
        Part("a", "production"), answer("NCL produced 12.70 MT [C1].", [shared]), reg, annexure_no=1
    )
    # Same chunk comes back as C2 in the next answer: it keeps its reply-wide number S1.
    same_again = Passage(
        "C2", shared.chunk_id, shared.document_id, shared.document_title, 1, None, ""
    )
    b2 = pq_reply.build_block(
        Part("b", "target"),
        answer("The target was 12.16 MT [C1] and output 12.70 MT [C2].", [other, same_again]),
        reg,
        annexure_no=1,
    )
    assert b1["text"] == "NCL produced 12.70 MT [S1]."
    assert b2["text"] == "The target was 12.16 MT [S2] and output 12.70 MT [S1]."
    assert b2["citations"] == ["S2", "S1"]
    assert [s["id"] for s in reg.sources] == ["S1", "S2"]
    assert reg.sources[0]["parts"] == ["a", "b"]
    assert reg.sources[1]["page"] == 33 and reg.sources[1]["kind"] == "chunk"


def test_part_without_source_says_not_found():
    reg = pq_reply.SourceRegistry()
    nf = answer(qa.NOT_FOUND, not_found=True)
    uncited = answer("Yes, the target was achieved.")  # words but no source at all
    for a in (nf, uncited):
        b = pq_reply.build_block(Part("b", "x"), a, reg, annexure_no=1)
        assert b["text"] == pq_reply.NOT_FOUND_PART
        assert b["not_found"] and b["citations"] == []
    hindi = pq_reply.build_block(Part("ख", "x"), nf, reg, annexure_no=1, hindi=True)
    assert hindi["text"] == pq_reply.NOT_FOUND_PART_HI
    assert reg.sources == []


def test_unsupported_figure_sentence_is_dropped():
    reg = pq_reply.SourceRegistry()
    a = answer(
        "NCL produced 12.70 MT [C1]. This was 9.4% higher than last year [C1].",
        [passage("C1")],
        unmatched=["9.4"],
    )
    b = pq_reply.build_block(Part("a", "x"), a, reg, annexure_no=1)
    assert b["text"] == "NCL produced 12.70 MT [S1]."
    assert b["removed_numbers"] == ["9.4"]

    only_bad = answer("Output rose by 9.4% [C1].", [passage("C1")], unmatched=["9.4"])
    b = pq_reply.build_block(Part("a", "x"), only_bad, reg, annexure_no=1)
    assert b["not_found"]


def test_small_table_stays_inline_large_table_goes_to_annexure():
    reg = pq_reply.SourceRegistry()
    few = [fact(f"F{i}") for i in range(1, 4)]
    small = {"columns": ["Company", "Value", "Source"], "rows": [[f.who, "1", f.ref] for f in few]}
    b = pq_reply.build_block(
        Part("a", "x"), answer("Figures [F1].", few[:1], table=small, facts=few), reg, annexure_no=1
    )
    assert b["annexure"] is None
    assert [r[2] for r in b["table"]["rows"]] == ["S1", "S2", "S3"]  # table refs renumbered
    assert b["citations"] == ["S1", "S2", "S3"]
    assert reg.sources[0]["kind"] == "fact" and reg.sources[0]["fact_id"]

    many = [fact(f"F{i}", who=f"CO{i}") for i in range(1, 9)]
    big = {"columns": ["Company", "Value", "Source"], "rows": [[f.who, "1", f.ref] for f in many]}
    text = "The coal production figures for Dec 2024 are given in the table below. [F1][F2]"
    b = pq_reply.build_block(
        Part("b", "x"), answer(text, many[:2], table=big, facts=many), reg, annexure_no=1
    )
    assert b["annexure"] == "Annexure-I"
    assert "are given in Annexure-I." in b["text"]
    assert len(b["table"]["rows"]) == 8


def test_annexures_are_numbered_and_links_point_at_facts_and_chunks():
    many = [fact(f"F{i}") for i in range(1, 9)]
    big = {"columns": ["C", "Source"], "rows": [[f.who, f.ref] for f in many]}
    answers = [
        answer("Figures [F1].", many[:1], table=big, facts=many),
        answer("NCL produced 12.70 MT [C1].", [passage("C1")]),
        answer("More figures [F1].", many[:1], table=big, facts=many),
    ]
    parts = [Part("a", "x"), Part("b", "y"), Part("c", "z")]
    draft = pq_reply.build_content(header(), parts, answers, hindi=False)
    blocks = draft.content["blocks"]
    assert [b["annexure"] for b in blocks] == ["Annexure-I", None, "Annexure-II"]
    assert "Details are given in Annexure-I." in blocks[0]["text"]
    chunk_links = [lk for lk in draft.links if lk["chunk_id"]]
    assert len(chunk_links) == 1 and chunk_links[0]["location_in_report"] == "part-b:S9"
    assert all((lk["fact_id"] is None) != (lk["chunk_id"] is None) for lk in draft.links)
    assert draft.links[0]["location_in_report"] == "Annexure-I:S1"
    assert len(draft.content["sources"]) == 9  # 8 facts shared by (a) and (c) + 1 chunk


def test_plain_text_rendering_has_header_parts_answers_and_sources():
    draft = pq_reply.build_content(
        header(),
        [Part("a", "the production?"), Part("b", "the target?")],
        [
            answer("NCL produced 12.70 MT [C1].", [passage("C1")]),
            answer(qa.NOT_FOUND, not_found=True),
        ],
        hindi=False,
    )
    text = draft.content["text"]
    for line in (
        "LOK SABHA",
        "UNSTARRED QUESTION NO. 1234",
        "TO BE ANSWERED ON 12.02.2025",
        "COAL PRODUCTION BY NCL",
        "1234. SHRI A. B. MEMBER:",
        "(a) the production?",
        "ANSWER",
        "(a): NCL produced 12.70 MT [S1].",
        f"(b): {pq_reply.NOT_FOUND_PART}",
        "[S1] Monthly Summary Dec 2024, page 1",
    ):
        assert line in text


def test_draft_reply_asks_the_qa_service_once_per_part():
    asked = []

    def fake_ask(db, user, question):
        asked.append(question)
        return answer("NCL produced 12.70 MT [C1].", [passage("C1")])

    draft = pq_reply.draft_reply(
        None,
        None,
        house=House.RAJYA_SABHA,
        q_type=QuestionType.STARRED,
        q_no="55",
        answer_date=None,
        mp_name=None,
        subject="NCL",
        question_text=PQ_EN,
        ask=fake_ask,
        llm=lambda *a, **k: "{}",
    )
    assert len(asked) == 3 and asked[0].startswith("NCL: the coal production")
    assert draft.split_method == "markers"
    assert draft.content["header"]["house_label"] == "RAJYA SABHA"
    assert "STARRED QUESTION NO. 55" in draft.content["text"]


# --- export -------------------------------------------------------------------------------------


def sample_report(status=ReportStatus.DRAFT):
    many = [fact(f"F{i}", who=f"CO{i}") for i in range(1, 9)]
    big = {"columns": ["Company", "Value", "Source"], "rows": [[f.who, "1", f.ref] for f in many]}
    draft = pq_reply.build_content(
        header(),
        [Part("a", "the production of NCL?"), Part("b", "subsidiary-wise?"), Part("c", "why?")],
        [
            answer("NCL produced 12.70 MT [C1].", [passage("C1")]),
            answer("Figures are given in the table below. [F1]", many[:1], table=big, facts=many),
            answer(qa.NOT_FOUND, not_found=True),
        ],
        hindi=False,
    )
    return SimpleNamespace(title="LS Q 1234", status=status, content_json=draft.content)


def docx_text(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in d.paragraphs] + cells)


def test_docx_export_contains_layout_parts_annexure_and_sources():
    view = report_export.view_from_report(sample_report(), None, [])
    text = docx_text(report_export.to_docx(view))
    for expected in (
        report_export.DRAFT_LINE,
        "MINISTRY OF COAL",
        "LOK SABHA",
        "UNSTARRED QUESTION NO. 1234",
        "TO BE ANSWERED ON 12.02.2025",
        "(a) the production of NCL?",
        "ANSWER",
        "(a): NCL produced 12.70 MT [S1].",
        "(b): Figures are given in Annexure-I.",
        f"(c): {pq_reply.NOT_FOUND_PART}",
        "ANNEXURE-I",
        "Referred to in reply to part (b) of Lok Sabha Unstarred Question No. 1234",
        "CO8",  # annexure table rows
        report_export.SOURCES_HEADING,
        "S1  Monthly Summary Dec 2024, page 1",
        "S2  Coal Directory, page 7",
    ):
        assert expected in text, expected


def test_approved_export_has_no_draft_line():
    view = report_export.view_from_report(sample_report(ReportStatus.APPROVED), None, [])
    assert report_export.DRAFT_LINE not in docx_text(report_export.to_docx(view))


def test_edited_plain_text_report_exports_with_pq_header_and_db_sources():
    report = SimpleNamespace(
        title="LS Q 1234", status=ReportStatus.IN_REVIEW, content_json={"text": "(a): Edited [S1]."}
    )
    sources = [{"id": "S1", "document": "Annual Report", "page": 3}]
    view = report_export.view_from_report(report, header(), sources)
    text = docx_text(report_export.to_docx(view))
    assert "UNSTARRED QUESTION NO. 1234" in text
    assert "(a): Edited [S1]." in text
    assert "S1  Annual Report, page 3" in text


def test_pdf_export_is_a_valid_pdf_with_the_reply():
    view = report_export.view_from_report(sample_report(), None, [])
    pdf = pymupdf.open("pdf", report_export.to_pdf(view))
    text = "".join(page.get_text() for page in pdf)
    assert "UNSTARRED QUESTION NO. 1234" in text
    assert "NCL produced 12.70 MT" in text
    assert "ANNEXURE-I" in text and "Coal Directory, page 7" in text


# --- endpoints ----------------------------------------------------------------------------------


class FakeSession:
    """Just enough of a Session for the draft endpoint: ids on flush, timestamps on refresh."""

    def __init__(self):
        self.added = []

    def add(self, obj):
        self.added.append(obj)

    def flush(self):
        for obj in self.added:
            if getattr(obj, "id", None) is None:
                obj.id = uuid.uuid4()

    def commit(self):
        self.flush()

    def refresh(self, obj):
        obj.created_at = obj.updated_at = datetime(2026, 9, 26)


PQ_BODY = {
    "house": "LS",
    "q_type": "unstarred",
    "q_no": "1234",
    "answer_date": "2025-02-12",
    "mp_name": "Shri A. B. Member",
    "subject": "Coal production by NCL",
    "question_text": PQ_EN,
}


def test_viewer_cannot_draft(client_as):
    assert client_as(["viewer"]).post("/pq-replies/draft", json=PQ_BODY).status_code == 403


def test_reviewer_cannot_export(client_as):
    rid = "00000000-0000-0000-0000-000000000001"
    assert client_as(["reviewer"]).get(f"/reports/{rid}/export").status_code == 403


def test_analyst_draft_saves_pq_report_versions_and_citations(client_as, monkeypatch):
    from app.core.db import get_db
    from app.main import app

    def fake_ask(db, user, question):
        if "details thereof" in question:
            return answer(qa.NOT_FOUND, not_found=True)
        return answer("NCL produced 12.70 MT [C1].", [passage("C1")])

    monkeypatch.setattr(qa, "ask", fake_ask)
    session = FakeSession()
    client = client_as(["analyst"])
    app.dependency_overrides[get_db] = lambda: session
    res = client.post("/pq-replies/draft", json=PQ_BODY)
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["split_method"] == "markers"
    report = body["report"]
    assert report["status"] == "DRAFT" and report["report_type"] == "pq_reply"
    assert report["title"] == "Lok Sabha Unstarred Q. No. 1234: Coal production by NCL"
    blocks = report["content_json"]["blocks"]
    assert [b["part"] for b in blocks] == ["a", "b", "c"]
    assert blocks[2]["text"] == pq_reply.NOT_FOUND_PART
    assert "text" in report["content_json"]

    kinds = [type(o) for o in session.added]
    assert kinds.count(ParliamentQuestion) == 1 and kinds.count(Report) == 1
    assert kinds.count(ReportVersion) == 1
    citations = [o for o in session.added if isinstance(o, ReportCitation)]
    assert len(citations) == 2  # parts (a) and (b), same kind of source each
    assert all(isinstance(c.chunk_id, uuid.UUID) and c.fact_id is None for c in citations)
    pq = next(o for o in session.added if isinstance(o, ParliamentQuestion))
    assert [p["label"] for p in pq.parts_json] == ["a", "b", "c"]
    audit_rows = [o for o in session.added if type(o).__name__ == "AuditLog"]
    assert audit_rows[0].action == "create_pq_draft"


def test_draft_when_model_is_down_is_503(client_as, monkeypatch):
    def down(*a, **k):
        raise client.LlmUnavailable("not reachable")

    monkeypatch.setattr(qa, "ask", down)
    res = client_as(["approver"]).post("/pq-replies/draft", json=PQ_BODY)
    assert res.status_code == 503
