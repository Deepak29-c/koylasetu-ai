"""Monthly report and briefing note first drafts (guide 10.1, W5, W6)."""

import io
import json
import uuid
from types import SimpleNamespace

import docx
import pymupdf
import pytest
from sqlalchemy.dialects import postgresql

from app.llm import client
from app.models import AuditLog, Report, ReportCitation, ReportVersion
from app.models.enums import ReportStatus, ReportType
from app.services import (
    briefing_note,
    monthly_report,
    pq_reply,
    qa,
    report_builder,
    report_export,
)
from app.services.report_builder import FactRow
from app.services.retrieval import Passage
from tests.test_pq_reply import FakeSession

DOC_A, DOC_B = uuid.uuid4(), uuid.uuid4()


def fact(metric, sub, period, value, status="VERIFIED", doc=DOC_A, page=1, confidence=1.0):
    return FactRow(
        id=uuid.uuid4(),
        metric=metric,
        subsidiary=sub,
        period=period,
        status=status,
        value_std=value,
        confidence=confidence,
        document_id=doc,
        document="Monthly Coal Statistics - December 2024" if doc == DOC_A else "Coal Directory",
        page=page,
    )


def dec_facts():
    """A small slice of the real December 2024 statistics (pages 1 and 3)."""
    return [
        fact("production_target", "NCL", "2024-12", 12.16),
        fact("coal_production", "NCL", "2024-12", 12.70),
        fact("coal_production", "NCL", "2023-12", 11.80),
        fact("production_target", "SECL", "2024-12", 18.30),
        fact("coal_production", "SECL", "2024-12", 14.33),
        fact("coal_production", "SECL", "2023-12", 16.60),
        fact("production_target", "CIL", "2024-12", 77.37),
        fact("coal_production", "CIL", "2024-12", 72.37),
        fact("coal_production", "CIL", "2023-12", 71.89),
        fact("coal_offtake", "CIL", "2024-12", 68.41, page=3),
        fact("coal_offtake", "CIL", "2023-12", 66.53, page=3),
    ]


def rows_by_code(facts, key="production"):
    spec = next(s for s in monthly_report.TABLES if s.key == key)
    return {r.code: r for r in monthly_report.build_rows(facts, "2024-12", spec)}


def table(content, key):
    return next(s for s in content["sections"] if s["id"] == key)["table"]


# --- the monthly table --------------------------------------------------------------------------


def test_achievement_and_growth_are_computed_by_code():
    rows = rows_by_code(dec_facts())
    assert rows["NCL"].achievement == 104.44  # 12.70 / 12.16
    assert rows["NCL"].growth == 7.63  # 12.70 / 11.80 - 1
    assert rows["SECL"].growth == -13.67
    assert monthly_report.pct(5, 0, change=False) is None
    assert monthly_report.pct(None, 5, change=True) is None


def test_missing_cells_show_a_dash_and_every_subsidiary_has_a_row():
    content = monthly_report.build_content(dec_facts(), "2024-12").content
    t = table(content, "production")
    assert t["columns"] == [
        "Subsidiary",
        "Target (MT)",
        "Actual (MT)",
        "Achievement (%)",
        "Dec 2023 (MT)",
        "Growth (%)",
        "Source",
    ]
    assert [r[0] for r in t["rows"]] == [label for _, label in monthly_report.ROWS]
    ecl = t["rows"][0]
    assert ecl == ["ECL", "–", "–", "–", "–", "–", "–"]
    ncl = next(r for r in t["rows"] if r[0] == "NCL")
    assert ncl == ["NCL", "12.16", "12.70", "104.44", "11.80", "7.63", "S1"]
    assert report_builder.NOT_FOUND_NOTE in content["text"]


def test_verified_fact_beats_a_more_confident_extracted_one_and_rejected_is_ignored():
    facts = [
        fact("coal_production", "NCL", "2024-12", 99.0, status="EXTRACTED", confidence=1.0),
        fact("coal_production", "NCL", "2024-12", 12.70, status="VERIFIED", confidence=0.6),
        fact("coal_production", "MCL", "2024-12", 50.0, status="REJECTED"),
        fact("coal_production", "MCL", "2024-12", 21.51, status="FLAGGED"),
        fact("coal_production", "WCL", "2024-12", 7.12, status="REJECTED"),
    ]
    rows = rows_by_code(facts)
    assert rows["NCL"].actual.value == 12.70
    assert rows["NCL"].actual.conflict  # the other document disagrees: shown as a warning
    assert rows["MCL"].actual.value == 21.51 and not rows["MCL"].actual.fact.verified
    assert rows["WCL"].actual.fact is None  # only a REJECTED fact: not found
    content = monthly_report.build_content(facts, "2024-12").content
    mcl = next(r for r in table(content, "production")["rows"] if r[0] == "MCL")
    assert mcl[2] == "21.51*"  # unverified figures are labelled
    assert report_builder.UNVERIFIED_NOTE in content["text"]
    assert any("NCL coal production" in w for w in content["warnings"])


def test_units_come_from_value_std_and_unconvertible_facts_are_skipped():
    facts = [
        fact("coal_production", "NCL", "2024-12", 12.7),  # e.g. 127 lakh tonnes -> 12.7 MT
        fact("coal_production", "MCL", "2024-12", None),  # unit not convertible to MT
    ]
    rows = rows_by_code(facts)
    assert rows["NCL"].actual.value == 12.7
    assert rows["MCL"].actual.fact is None


def test_citations_point_at_each_fact_and_sources_are_numbered_by_page():
    facts = dec_facts() + [fact("coal_production", "MCL", "2024-12", 21.51, doc=DOC_B, page=7)]
    draft = monthly_report.build_content(facts, "2024-12")
    sources = draft.content["sources"]
    assert [(s["id"], s["page"]) for s in sources] == [("S1", 1), ("S2", 7), ("S3", 3)]
    assert len(draft.links) == len(facts)  # one report_citations row per fact used
    assert all(link["fact_id"] and link["chunk_id"] is None for link in draft.links)
    assert "production:NCL:actual:S1" in {link["location_in_report"] for link in draft.links}
    rebuilt = pq_reply.sources_from_links(
        [(ln["location_in_report"], ln["fact_id"], "doc", 1) for ln in draft.links]
    )
    assert [s["id"] for s in rebuilt] == ["S1", "S2", "S3"]  # shared numbers listed once


def test_summary_lines_are_written_by_code_with_sources():
    content = monthly_report.build_content(dec_facts(), "2024-12").content
    summary = next(s for s in content["sections"] if s["id"] == "summary")["lines"]
    assert summary[0] == (
        "CIL coal production in December 2024 was 72.37 MT against a target of 77.37 MT "
        "(93.54% achievement), 0.67% higher than 71.89 MT in December 2023 [S1]."
    )
    assert summary[1].startswith("CIL coal despatch in December 2024 was 68.41 MT")


# --- remarks: the only model-written words -------------------------------------------------------


def test_model_remarks_are_grounded_against_the_rows_they_cite():
    prompts = []

    def llm(system, user, **kw):
        prompts.append(user)
        return (
            "SECL production fell 13.67% below December 2023 [F2]. "
            "NCL exceeded its target at 104.44% [F1]. "
            "CIL production rose 9.9% [F3]. "  # not in row 3: dropped
            "Overall performance was satisfactory."  # no figure, no citation: dropped
        )

    content = monthly_report.build_content(dec_facts(), "2024-12", llm=llm).content
    remarks = next(s for s in content["sections"] if s["id"] == "remarks")
    assert remarks["lines"] == [
        "SECL production fell 13.67% below December 2023 [S1].",
        "NCL exceeded its target at 104.44% [S1].",
    ]
    assert remarks["removed_numbers"] == ["9.9"]
    assert content["remarks_by"] == "model"
    assert "[F1] NCL coal production, December 2024: actual 12.70 MT;" in prompts[0]


def test_model_down_falls_back_to_rule_remarks_with_a_warning():
    def down(*a, **k):
        raise client.LlmUnavailable("no server")

    content = monthly_report.build_content(dec_facts(), "2024-12", llm=down).content
    remarks = next(s for s in content["sections"] if s["id"] == "remarks")["lines"]
    assert remarks[0] == (
        "SECL achieved 78.31% of its coal production target (14.33 MT against 18.30 MT) [S1]."
    )
    assert any("SECL coal production was 13.67% lower" in r for r in remarks)
    assert content["remarks_by"] == "rule"
    assert any("not available" in w for w in content["warnings"])


def test_bad_month_is_refused():
    with pytest.raises(ValueError):
        monthly_report.build_content([], "2024-13")


# --- briefing note --------------------------------------------------------------------------------


def test_topic_words_pick_metrics_and_subsidiaries():
    assert briefing_note.topic_metrics("Coking coal production status") == [
        "coking_coal_production"
    ]
    assert briefing_note.topic_metrics("NCL despatch") == ["coal_offtake"]
    assert briefing_note.topic_metrics("Meeting with the Minister") == [
        "coal_production",
        "coal_offtake",
    ]
    assert briefing_note.topic_subsidiaries("NCL and MCL production") == ["NCL", "MCL"]
    assert briefing_note.topic_subsidiaries("India's coal output") == ["INDIA"]
    assert briefing_note.topic_subsidiaries("coal output") == ["INDIA", "CIL"]


def test_periods_in_range():
    assert briefing_note.in_range("2024-12", "2024-04", "2024-12")
    assert not briefing_note.in_range("2025-01", "2024-04", "2024-12")
    assert briefing_note.in_range("FY 2024-25", "2025-01", None)  # overlaps
    assert not briefing_note.in_range("FY 2022-23", "2024-01", None)
    assert briefing_note.year_before("FY 2024-25") == "FY 2023-24"


def passage(ref, text, title="PIB Year End Review 2024", page=3):
    return Passage(ref, uuid.uuid4(), uuid.uuid4(), title, page, None, text)


def ans(text, passages=(), not_found=False):
    return SimpleNamespace(
        answer=text, citations=list(passages), passages=list(passages), not_found=not_found
    )


def briefing_facts():
    return [
        fact("coal_production", "CIL", "2024-11", 70.0),
        fact("coal_production", "CIL", "2024-12", 72.37),
        fact("coal_production", "CIL", "2023-12", 71.89),
        fact("coal_production", "CIL", "FY 2023-24", 773.81, status="EXTRACTED", page=5),
        fact("coal_production", "INDIA", "2024-12", 97.87, page=2),
    ]


def test_briefing_assembles_figures_trend_points_and_sources():
    p1 = passage("C1", "Coal production reached 1047.57 MT in 2024-25.")
    p2 = passage("C2", "Commercial coal mining auctions: 10th round launched.", page=9)
    points = ans(
        "India's coal production reached 1047.57 MT [C1]. It rose 12% [C1]. "
        "Auctions continue [C2].",
        [p1, p2],
    )
    draft = briefing_note.build_content(
        "India coal production", briefing_facts(), points, ans(qa.NOT_FOUND, not_found=True)
    )
    c = draft.content
    assert c["subsidiaries"] == ["INDIA"] and c["metrics"] == ["coal_production"]
    key = table(c, "key_figures")
    assert key["rows"] == [
        ["Coal production - All India (MT)", "Dec 2024", "97.87", "–", "–", "S1"]
    ]
    kp = next(s for s in c["sections"] if s["id"] == "key_points")
    assert kp["bullets"] and kp["lines"] == [
        "India's coal production reached 1047.57 MT [S2].",
        "Auctions continue [S3].",
    ]
    assert kp["removed_numbers"] == ["12"]
    steps = next(s for s in c["sections"] if s["id"] == "steps")
    assert steps["lines"] == [briefing_note.NOT_FOUND_POINT]
    assert [s["page"] for s in c["sources"]] == [2, 3, 9]
    chunk_links = [ln for ln in draft.links if ln["chunk_id"]]
    assert {ln["location_in_report"] for ln in chunk_links} == {"key_points:S2", "key_points:S3"}
    assert "• India's coal production reached 1047.57 MT [S2]." in c["text"]


def test_briefing_key_figures_have_change_and_trend_from_facts():
    c = briefing_note.build_content(
        "CIL production", briefing_facts(), ans(qa.NOT_FOUND, not_found=True), ans("x")
    ).content
    rows = table(c, "key_figures")["rows"]
    assert rows[0] == ["Coal production - CIL (MT)", "Dec 2024", "72.37", "71.89", "0.67", "S1"]
    assert rows[1][:3] == ["Coal production - CIL (MT)", "FY 2023-24", "773.81*"]
    trend = table(c, "trend")
    assert [r[:2] for r in trend["rows"]] == [
        ["Dec 2023", "71.89"],
        ["Nov 2024", "70.00"],
        ["Dec 2024", "72.37"],
    ]
    chart = report_builder.chart_for_table(trend, "TREND")
    assert chart["labels"] == ["Dec 2023", "Nov 2024", "Dec 2024"]
    ranged = briefing_note.build_content(
        "CIL production", briefing_facts(), ans("x"), ans("x"), start="2024-11", end="2024-12"
    ).content
    assert [r[0] for r in table(ranged, "trend")["rows"]] == ["Nov 2024", "Dec 2024"]


def test_draft_briefing_asks_qa_twice_and_checks_the_model_first(monkeypatch):
    asked = []
    monkeypatch.setattr(briefing_note, "fetch_facts", lambda *a, **k: briefing_facts())

    def ask(db, user, question):
        asked.append(question)
        return ans(qa.NOT_FOUND, not_found=True)

    briefing_note.draft_briefing(None, None, "CIL production", ask=ask, check_model=lambda: None)
    assert asked == [
        "CIL production: key facts, figures and current status",
        "CIL production: recent steps and initiatives taken",
    ]

    def down():
        raise client.LlmUnavailable("down")

    asked.clear()
    with pytest.raises(client.LlmUnavailable):
        briefing_note.draft_briefing(None, None, "x", ask=ask, check_model=down)
    assert asked == []


# --- export ---------------------------------------------------------------------------------------


def report(content, status=ReportStatus.DRAFT, report_type=ReportType.MONTHLY):
    return SimpleNamespace(
        title="Monthly report - December 2024",
        status=status,
        report_type=report_type,
        content_json=content,
    )


def docx_parts(data: bytes):
    d = docx.Document(io.BytesIO(data))
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    return "\n".join([p.text for p in d.paragraphs] + cells), d


def test_monthly_docx_has_table_chart_remarks_and_sources():
    content = monthly_report.build_content(dec_facts(), "2024-12").content
    items = report_export.items_for_report(report(content), None, [])
    text, d = docx_parts(report_export.to_docx(items))
    for expected in (
        pq_reply.DRAFT_LINE,
        monthly_report.TITLE,
        "December 2024 (figures in million tonnes, MT)",
        "SUMMARY",
        "1. COAL PRODUCTION",
        "104.44",
        "REMARKS",
        "SOURCES",
        "[S1] Monthly Coal Statistics - December 2024, page 1",
    ):
        assert expected in text, expected
    assert len(d.tables) == 3
    # One chart: production. Despatch has only CIL here (totals are not charted) and coking
    # has no figures, so there is nothing to draw for them.
    assert len(d.inline_shapes) == 1


def test_approved_monthly_export_has_no_draft_line():
    content = monthly_report.build_content(dec_facts(), "2024-12").content
    items = report_export.items_for_report(report(content, ReportStatus.APPROVED), None, [])
    assert pq_reply.DRAFT_LINE not in docx_parts(report_export.to_docx(items))[0]


def test_monthly_saved_by_the_plain_editor_keeps_its_table_chart_and_sources():
    content = monthly_report.build_content(dec_facts(), "2024-12").content
    saved = {"text": content["text"].replace("12.70", "12.71")}  # an edit in the editor
    db_sources = [{"id": "S1", "document": "Monthly Coal Statistics - December 2024", "page": 1}]
    items = report_export.items_for_report(report(saved), None, db_sources)
    text, d = docx_parts(report_export.to_docx(items))
    assert text.count("SOURCES") == 1 and text.count(monthly_report.TITLE) == 1
    assert "12.71" in text
    assert len(d.tables) == 3 and len(d.inline_shapes) == 1
    chart = next(i[1] for i in items if i[0] == "chart")
    assert chart["title"] == "1. COAL PRODUCTION"
    ncl = chart["labels"].index("NCL")
    assert chart["series"][1] == {"name": "Actual", "values": chart["series"][1]["values"]}
    assert chart["series"][1]["values"][ncl] == 12.71  # the chart follows the edited table
    assert "CIL" not in chart["labels"] and "All India" not in chart["labels"]


def test_monthly_pdf_contains_the_table_and_a_chart_image():
    content = monthly_report.build_content(dec_facts(), "2024-12").content
    items = report_export.items_for_report(report(content), None, [])
    pdf = pymupdf.open("pdf", report_export.to_pdf(items))
    text = "".join(page.get_text() for page in pdf)
    assert "1. COAL PRODUCTION" in text and "104.44" in text
    assert "Monthly Coal Statistics - December 2024, page 1" in text
    assert sum(len(page.get_images()) for page in pdf) == 1


def test_briefing_pdf_renders_hindi_points():
    p1 = passage("C1", "कोयला उत्पादन 1047.57 मिलियन टन रहा।")
    c = briefing_note.build_content(
        "कोयला उत्पादन", briefing_facts(), ans("कोयला उत्पादन 1047.57 मिलियन टन रहा [C1]।", [p1]),
        ans(qa.NOT_FOUND, not_found=True),
    ).content  # fmt: skip
    items = report_export.items_for_report(report(c, report_type=ReportType.BRIEFING), None, [])
    text = "".join(p.get_text() for p in pymupdf.open("pdf", report_export.to_pdf(items)))
    assert "1047.57" in text and "BRIEFING NOTE" in text and "KEY POINTS" in text


def test_axis_ticks_are_round_numbers():
    assert report_export.axis_ticks(21.51) == [0, 5, 10, 15, 20, 25]
    assert report_export.axis_ticks(0.58) == [0, 0.2, 0.4, 0.6]
    assert report_export.chart_png(
        {"title": "t", "unit": "MT", "labels": ["A"], "series": [{"name": "x", "values": [1]}]}
    ).startswith(b"\x89PNG")


# --- the facts query ------------------------------------------------------------------------------


class CapturingDb:
    def __init__(self):
        self.stmt = None

    def execute(self, stmt):
        self.stmt = stmt
        return SimpleNamespace(all=lambda: [])


def sql(stmt) -> str:
    return str(stmt.compile(dialect=postgresql.dialect()))


def test_fact_query_is_scoped_through_the_document_and_leaves_out_rejected():
    db = CapturingDb()
    unit = SimpleNamespace(org_unit=SimpleNamespace(code="NCL"), org_unit_id=uuid.uuid4())
    report_builder.fetch_facts(db, unit, metrics=["coal_production"], periods=["2024-12"])
    where = sql(db.stmt)
    assert "documents.org_unit_id" in where and "documents.is_public" in where
    assert "facts.status !=" in where and "facts.period IN" in where
    hq = SimpleNamespace(org_unit=SimpleNamespace(code="CIL_HQ"), org_unit_id=uuid.uuid4())
    report_builder.fetch_facts(db, hq, metrics=["coal_production"])
    assert "org_unit_id" not in sql(db.stmt)  # CIL HQ sees every unit


# --- endpoints ------------------------------------------------------------------------------------


@pytest.fixture
def analyst(client_as, monkeypatch):
    from app.core.db import get_db
    from app.main import app

    session = FakeSession()
    http = client_as(["analyst"])
    app.dependency_overrides[get_db] = lambda: session
    monkeypatch.setattr(monthly_report, "fetch_facts", lambda *a, **k: dec_facts())
    monkeypatch.setattr(client, "chat", lambda *a, **k: "NCL achieved 104.44% [F1].")
    return http, session


def test_viewer_cannot_draft_monthly_or_briefing(client_as):
    viewer = client_as(["viewer"])
    assert viewer.post("/reports/monthly/draft", json={"month": "2024-12"}).status_code == 403
    assert viewer.post("/reports/briefing/draft", json={"topic": "coal"}).status_code == 403


def test_bad_input_is_422(analyst):
    http, _ = analyst
    assert http.post("/reports/monthly/draft", json={"month": "Dec 2024"}).status_code == 422
    body = {"month": "2024-12", "subsidiaries": ["XYZ"]}
    assert http.post("/reports/monthly/draft", json=body).status_code == 422
    body = {"topic": "coal", "from": "2024-12", "to": "2024-01"}
    assert http.post("/reports/briefing/draft", json=body).status_code == 422


def test_monthly_draft_saves_report_version_citations_and_audit(analyst):
    http, session = analyst
    res = http.post("/reports/monthly/draft", json={"month": "2024-12"})
    assert res.status_code == 201, res.text
    r = res.json()["report"]
    assert r["report_type"] == "monthly" and r["status"] == "DRAFT"
    assert r["title"] == "Monthly report - December 2024"
    assert r["content_json"]["text"].startswith(monthly_report.TITLE)
    remarks = next(s for s in r["content_json"]["sections"] if s["id"] == "remarks")
    assert remarks["lines"] == ["NCL achieved 104.44% [S1]."]
    kinds = [type(o) for o in session.added]
    assert kinds.count(Report) == 1 and kinds.count(ReportVersion) == 1
    citations = [o for o in session.added if isinstance(o, ReportCitation)]
    assert len(citations) == len(dec_facts())
    assert all(c.location_in_report.rsplit(":", 1)[1].startswith("S") for c in citations)
    audit_row = next(o for o in session.added if isinstance(o, AuditLog))
    assert audit_row.action == "create_monthly_draft"
    assert audit_row.details_json["facts_used"] == len(dec_facts())
    json.dumps(audit_row.details_json)


def test_monthly_retry_with_the_same_key_returns_the_first_draft(analyst):
    http, session = analyst
    key = {"Idempotency-Key": str(uuid.uuid4())}
    first = http.post("/reports/monthly/draft", json={"month": "2024-12"}, headers=key)
    second = http.post("/reports/monthly/draft", json={"month": "2024-12"}, headers=key)
    assert first.status_code == 201 and second.status_code == 200
    assert first.json()["report"]["id"] == second.json()["report"]["id"]
    assert [type(o) for o in session.added].count(Report) == 1


def test_briefing_draft_uses_the_qa_service(analyst, monkeypatch):
    http, session = analyst
    monkeypatch.setattr(client, "ensure_available", lambda: None)
    monkeypatch.setattr(briefing_note, "fetch_facts", lambda *a, **k: briefing_facts())
    p1 = passage("C1", "CIL produced 781.06 MT in 2024-25.")
    monkeypatch.setattr(qa, "ask", lambda db, user, q: ans("CIL produced 781.06 MT [C1].", [p1]))
    res = http.post("/reports/briefing/draft", json={"topic": "CIL production", "to": "2024-12"})
    assert res.status_code == 201, res.text
    r = res.json()["report"]
    assert r["report_type"] == "briefing" and r["title"] == "Briefing note: CIL production"
    assert "• CIL produced 781.06 MT [S" in r["content_json"]["text"]
    audit_row = next(o for o in session.added if isinstance(o, AuditLog))
    assert audit_row.action == "create_briefing_draft"
    citations = [o for o in session.added if isinstance(o, ReportCitation)]
    assert any(c.chunk_id for c in citations) and any(c.fact_id for c in citations)


def test_briefing_when_model_is_down_is_503_and_saves_nothing(analyst, monkeypatch):
    http, session = analyst

    def down():
        raise client.LlmUnavailable("model server not reachable at http://10.0.0.5:11434")

    monkeypatch.setattr(client, "ensure_available", down)
    res = http.post("/reports/briefing/draft", json={"topic": "CIL production"})
    assert res.status_code == 503 and "10.0.0.5" not in res.json()["detail"]
    assert session.added == []
