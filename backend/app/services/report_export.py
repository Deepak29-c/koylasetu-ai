"""DOCX and PDF export of a report in the official PQ-reply layout (guide 10.1).

The layout itself (Ministry header, house, Starred/Unstarred Question No., date, subject, MP
and the parts, ANSWER, Minister line placeholder, Statement for starred questions, answers,
Annexures, internal Sources list) is built once in pq_reply.layout as lines and tables; this
module only draws those items.

DOCX is built with python-docx and PDF with PyMuPDF's Story (HTML -> PDF), both already
dependencies, so no LibreOffice or new package is needed. PyMuPDF shapes Devanagari, so Hindi
replies render correctly. A report that is not APPROVED carries a "DRAFT" line, its warnings
and a note under each part where unsupported figures were removed.

Two sources of content: a generated reply (blocks in content_json), or a reply edited in the
plain-text editor, which saves only {text}. That text is already the full reply (render_text
output, edited), so it is drawn as it is; the header and sources are only added when the text
does not have them.
"""

from __future__ import annotations

import html
import io
import re

import pymupdf
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from app.services.pq_reply import (
    DRAFT_LINE,
    SOURCES_HEADING,
    header_lines,
    layout,
)

_MARK = re.compile(r"(\[S\d+\])")
_CENTRED = re.compile(r"^(ANSWER|STATEMENT|ANNEXURE-[IVXL]+|THE MINISTER OF .*)$")
_CAPTION = re.compile(r"^(Statement referred to|Referred to) in reply to ")
_ASKER = re.compile(r"^[\w*/-]+\. [^a-z]+:$")  # "1234. SHRI A. B. MEMBER:"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


def is_approved(report) -> bool:
    return report.status.value in ("APPROVED", "EXPORTED")


def items_from_text(
    text: str, *, header: dict | None, sources: list[dict], title: str, approved: bool
) -> list[tuple]:
    """Layout items for a reply edited as plain text. Lines up to the first blank line are the
    header when the text starts with it; "a | b" line runs become tables."""
    items: list[tuple] = []
    if not approved:
        items.append(("line", DRAFT_LINE, "note"))
    lines = text.strip("\n").splitlines()
    has_header = bool(lines) and lines[0].strip() == "GOVERNMENT OF INDIA"
    if not has_header:
        if header:
            items += [("line", h, "title") for h in header_lines(header)] + [("line", "", "normal")]
        else:
            items.append(("line", title, "title"))
    in_header, in_sources, table_rows = has_header, False, []

    def flush_table() -> None:
        if table_rows:
            items.append(("table", {"columns": table_rows[0], "rows": table_rows[1:]}))
            table_rows.clear()

    subject_next = False
    for raw in lines:
        line = raw.strip()
        if " | " in line and not in_sources:
            table_rows.append([c.strip() for c in line.split(" | ")])
            continue
        flush_table()
        if not line:
            subject_next = in_header  # render_text puts the subject after the header block
            in_header = False
            items.append(("line", "", "normal"))
            continue
        if line == SOURCES_HEADING:
            in_sources = True
            items.append(("line", line, "bold"))
        elif in_sources:
            items.append(("line", line, "small"))
        elif in_header or subject_next or _CENTRED.match(line):
            items.append(("line", line, "title"))
        elif _CAPTION.match(line):
            items.append(("line", line, "caption"))
        elif _ASKER.match(line):
            items.append(("line", line, "bold"))
        else:
            items.append(("line", line, "normal"))
        subject_next = False
    flush_table()
    if sources and SOURCES_HEADING not in text:
        items += [("line", "", "normal"), ("line", SOURCES_HEADING, "bold")]
        items += [
            ("line", f"[{s['id']}] {s['document']}, page {s['page']}", "small") for s in sources
        ]
    return items


def items_for_report(report, pq_header: dict | None, db_sources: list[dict]) -> list[tuple]:
    content = report.content_json or {}
    approved = is_approved(report)
    blocks = content.get("blocks") or []
    header = content.get("header") or pq_header
    if blocks and header:
        return layout(
            header,
            blocks,
            content.get("sources") or db_sources,
            internal_notes=not approved,
            warnings=content.get("warnings") or [],
        )
    return items_from_text(
        content.get("text", ""),
        header=pq_header,
        sources=db_sources,
        title=report.title,
        approved=approved,
    )


# --- DOCX ---------------------------------------------------------------------------------------


def _para(doc, text: str, style: str):
    p = doc.add_paragraph()
    if style in ("title", "caption"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pieces = [text] if style == "small" else _MARK.split(text)
    for piece in pieces:
        if not piece:
            continue
        run = p.add_run(piece)
        run.bold = style in ("title", "bold")
        if style != "small" and _MARK.fullmatch(piece):  # source markers: like footnote numbers
            run.font.superscript = True
            run.font.size = Pt(8)
        elif style in ("caption", "note"):
            run.font.size = Pt(9.5)
        elif style == "small":
            run.font.size = Pt(9)
        if style == "note":
            run.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)
    return p


def _docx_table(doc, table: dict) -> None:
    width = max([len(table["columns"])] + [len(r) for r in table["rows"]])
    t = doc.add_table(rows=1, cols=width)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, table["columns"], strict=False):
        cell.text = ""
        cell.paragraphs[0].add_run(str(text)).bold = True
    for row in table["rows"]:
        for cell, text in zip(t.add_row().cells, row, strict=False):
            cell.text = str(text)
    doc.add_paragraph()


def to_docx(items: list[tuple]) -> bytes:
    doc = DocxDocument()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    for item in items:
        if item[0] == "line":
            _para(doc, item[1], item[2])
        elif item[0] == "table":
            _docx_table(doc, item[1])
        else:
            doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF ----------------------------------------------------------------------------------------

_CSS = """
body { font-family: sans-serif; font-size: 11pt; }
p { margin: 0 0 5pt 0; }
.title { text-align: center; font-weight: bold; }
.bold { font-weight: bold; }
.caption { text-align: center; font-size: 10pt; }
.note { color: #b00000; font-size: 9pt; }
.small { font-size: 9pt; margin: 0; }
sup { font-size: 7pt; }
table { border-collapse: collapse; margin: 4pt 0 8pt 0; }
td, th { border: 1px solid #000; padding: 2pt 4pt; font-size: 9.5pt; }
"""


def _html_table(table: dict) -> str:
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in table["columns"])
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in row) + "</tr>"
        for row in table["rows"]
    )
    return f"<table><tr>{head}</tr>{body}</table>"


def to_html(items: list[tuple]) -> str:
    out: list[str] = []
    for item in items:
        if item[0] == "line":
            text, style = item[1], item[2]
            body = html.escape(text)
            if style != "small":
                body = _MARK.sub(lambda m: f"<sup>{m.group(1)}</sup>", body)
            out.append(f'<p class="{style}">{body or "&#160;"}</p>')
        elif item[0] == "table":
            out.append(_html_table(item[1]))
        else:
            out.append("<br/>")
    return "<body>" + "".join(out) + "</body>"


def to_pdf(items: list[tuple]) -> bytes:
    story = pymupdf.Story(html=to_html(items), user_css=_CSS)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    page = pymupdf.paper_rect("a4")
    area = page + (60, 56, -60, -56)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buf.getvalue()
